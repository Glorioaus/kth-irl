"""AG1 M2 合成材料测试；不使用真实材料、模型或研究服务。"""

from __future__ import annotations

import copy
import io
import json
import sqlite3
import struct
import sys
import zipfile
from pathlib import Path

import docx
import pytest
from PyPDF2 import PdfWriter
from PyPDF2.generic import (
    ArrayObject, DecodedStreamObject, DictionaryObject, NameObject,
)

from kth_hybrid import intake
from kth_hybrid.agent_host import ProductRejected, digest
from kth_hybrid.contracts import sha256_hex
from kth_hybrid.workflow import LocalWorkflow


INPUT_FIELDS = {
    "input_index", "origin_path", "original_filename", "blob_sha256",
    "byte_length", "status", "error",
}


@pytest.fixture
def workflow(tmp_path):
    instance = LocalWorkflow(tmp_path / "case")
    yield instance
    instance.close()


def _api():
    freeze = getattr(intake, "freeze_product_inputs", None)
    project = getattr(intake, "project_product_materials", None)
    assert callable(freeze), "M2缺少冻结输入能力"
    assert callable(project), "M2缺少完整材料覆盖能力"
    return freeze, project


def _project(workflow, paths, run_id="RUN::synthetic"):
    freeze, project = _api()
    inputs = freeze([str(path) for path in paths], workflow.blobs)
    return inputs, project(workflow, run_id, inputs)


def _text(workflow, segment):
    data = workflow.blobs.read_bytes(segment["text_blob_sha256"])
    assert sha256_hex(data) == segment["text_sha256"]
    text = data.decode("utf-8")
    assert len(text) == segment["char_count"]
    return text


def _pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    page = writer.pages[0]
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        NameObject("/ProcSet"): ArrayObject([NameObject("/PDF"), NameObject("/Text")]),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 100 Td (synthetic PDF evidence) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_freeze_exact_fields_and_missing_denominator(workflow, tmp_path):
    path = tmp_path / "input.txt"
    path.write_bytes("合成输入\r\n".encode("utf-8"))
    before = path.stat()
    inputs, coverage = _project(workflow, [path, tmp_path / "missing.txt", tmp_path])
    assert len(inputs) == coverage["received_count"] == 3
    assert all(set(item) == INPUT_FIELDS for item in inputs)
    assert [item["input_index"] for item in inputs] == [0, 1, 2]
    assert inputs[0]["status"] == "frozen", inputs[0]["error"]
    assert inputs[0]["byte_length"] == 14
    assert inputs[0]["blob_sha256"] == sha256_hex("合成输入\r\n".encode("utf-8"))
    assert all(item["blob_sha256"] is None and item["error"] for item in inputs[1:])
    assert len(coverage["entries"]) == 3
    assert len(coverage["unprocessed"]) == 2
    assert path.stat().st_mtime_ns == before.st_mtime_ns
    assert _text(workflow, coverage["segments"][0]) == "合成输入\r\n"


def test_projection_uses_frozen_bytes_and_is_persisted_idempotently(workflow, tmp_path):
    path = tmp_path / "original.txt"
    path.write_bytes(b"frozen original\n")
    freeze, project = _api()
    inputs = freeze([str(path)], workflow.blobs)
    path.write_bytes(b"replacement")
    coverage = project(workflow, "RUN::frozen", inputs)
    path.unlink()
    before = {table: len(workflow.store.fetch_all(table))
              for table in ("sources", "import_records")}
    projections = workflow.store.fetch_text_projections(coverage["entries"][0]["source_id"])
    assert project(workflow, "RUN::frozen", inputs) == coverage
    assert {table: len(workflow.store.fetch_all(table)) for table in before} == before
    assert workflow.store.fetch_text_projections(
        coverage["entries"][0]["source_id"]) == projections
    assert _text(workflow, coverage["segments"][0]) == "frozen original\n"
    assert coverage["schema_version"] == "ag1.coverage.v1"
    assert coverage["run_id"] == "RUN::frozen"
    assert coverage["coverage_digest"] == digest({
        key: value for key, value in coverage.items() if key != "coverage_digest"})
    index = workflow.store.get_product_object(
        "COVERAGE::RUN::frozen", run_id="RUN::frozen", kind="coverage")
    assert index["schema_version"] == "ag1.coverage.index.v1"
    assert index["coverage_digest"] == coverage["coverage_digest"]
    assert index["counts"] == {"entries": 1, "segments": 1, "unprocessed": 0}
    assert not {"entries", "segments", "unprocessed"} & set(index)
    assert index["parts"]
    assert all("text" not in item for item in coverage["segments"])


def test_same_blob_different_occurrences_have_distinct_sources(workflow, tmp_path):
    path = tmp_path / "same.txt"
    other = tmp_path / "other.md"
    path.write_bytes(b"same content\n")
    other.write_bytes(path.read_bytes())
    inputs, coverage = _project(workflow, [path, path, other])
    entries = coverage["entries"]
    assert len({item["material_id"] for item in entries}) == 3
    assert len({item["source_id"] for item in entries}) == 3
    assert len({item["blob_sha256"] for item in entries}) == 1
    assert len(workflow.store.fetch_all("import_records")) == 3
    for entry in entries:
        source = workflow.store.fetch_one("sources", "source_id", entry["source_id"])
        assert source["published_at"] is None
        assert source["retrieved_at"] is None
        assert source["document_subject"] is None
        assert workflow.store.fetch_text_projections(entry["source_id"])
    _, second = _project(workflow, [path], "RUN::different")
    assert second["entries"][0]["source_id"] not in {e["source_id"] for e in entries}
    changed = copy.deepcopy(inputs)
    changed[0]["origin_path"] += ".changed"
    with pytest.raises((intake.IntakeRejected, ValueError), match="身份|输入|冲突"):
        intake.project_product_materials(workflow, "RUN::synthetic", changed)


def test_pdf_text_page_and_blank_page_are_both_covered(workflow, tmp_path):
    path = tmp_path / "pages.pdf"
    path.write_bytes(_pdf_bytes())
    _, coverage = _project(workflow, [path])
    assert len(coverage["segments"]) == 1
    segment = coverage["segments"][0]
    assert segment["locator"]["page"] == 1
    assert "synthetic PDF evidence" in _text(workflow, segment)
    assert segment["tool"].startswith("PyPDF2")
    assert any(item["locator"].get("page") == 2 for item in coverage["unprocessed"])
    assert coverage["entries"][0]["status"] == "partial"


def test_docx_tables_and_empty_paragraphs_have_locators(workflow, tmp_path):
    document = docx.Document()
    document.add_paragraph("第一段")
    document.add_paragraph("")
    document.add_paragraph("第三段")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "技术"
    table.cell(0, 1).text = "验证"
    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "未完成"
    path = tmp_path / "table.docx"
    document.save(path)
    _, coverage = _project(workflow, [path])
    assert {_text(workflow, s) for s in coverage["segments"]} == {
        "第一段", "第三段", "技术", "验证", "未完成"}
    cell = next(s for s in coverage["segments"] if _text(workflow, s) == "未完成")
    assert cell["locator"] == {"table": 1, "row": 2, "column": 2}
    assert any(u["locator"] == {"paragraph": 2} for u in coverage["unprocessed"])
    assert any(u["locator"] == {"table": 1, "row": 2, "column": 1}
               for u in coverage["unprocessed"])
    legacy = intake.extract_docx_paragraphs(path.read_bytes())
    assert [item["paragraph"] for item in legacy.locators] == [1, 3]
    assert all("paragraph" in item for item in legacy.locators)


def _write_docx_stories(path):
    """自有标准OOXML部件，不读取python-docx的页眉/页脚模板。"""
    from xml.etree import ElementTree as ET

    document = docx.Document()
    document.add_paragraph("BODY_FIRST")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "BODY_CELL"
    document.add_paragraph("BODY_LAST")
    original = io.BytesIO()
    document.save(original)
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    relationships = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_types = "http://schemas.openxmlformats.org/package/2006/content-types"
    body = ET.fromstring(members["word/document.xml"])
    section = body.find(f".//{{{w}}}sectPr")
    rels = ET.fromstring(members["word/_rels/document.xml.rels"])
    types = ET.fromstring(members["[Content_Types].xml"])
    for story, tag, label in (("header", "hdr", "HEADER"), ("footer", "ftr", "FOOTER")):
        member = f"word/{story}1.xml"
        relationship_id = f"rIdSynthetic{label}"
        ET.SubElement(section, f"{{{w}}}{story}Reference", {
            f"{{{w}}}type": "default", f"{{{r}}}id": relationship_id})
        ET.SubElement(rels, f"{{{relationships}}}Relationship", {
            "Id": relationship_id, "Type": f"{r}/{story}", "Target": f"{story}1.xml"})
        ET.SubElement(types, f"{{{content_types}}}Override", {
            "PartName": "/" + member,
            "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{story}+xml"})
        members[member] = (
            f'<w:{tag} xmlns:w="{w}">'
            f'<w:p><w:r><w:t>{label}_VISIBLE</w:t></w:r></w:p>'
            '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="3000"/></w:tblGrid>'
            '<w:tr><w:tc><w:tcPr/><w:p><w:r>'
            f'<w:t>{label}_CELL</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
            f'</w:{tag}>').encode("utf-8")
    members["word/document.xml"] = ET.tostring(body, encoding="utf-8")
    members["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8")
    members["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def _write_docx_altchunk_in_table(path, *, nested=False, external=False):
    from xml.etree import ElementTree as ET

    _write_docx_stories(path)
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    document = ET.fromstring(members["word/document.xml"])
    table = document.find(f".//{{{w}}}tbl")
    cell = table.find(f".//{{{w}}}tc")
    altchunk = ET.Element(f"{{{w}}}altChunk", {f"{{{r}}}id": "chunk"})
    if not nested:
        cell.insert(2, altchunk)
    else:
        nested_table = ET.Element(f"{{{w}}}tbl")
        nested_row = ET.SubElement(nested_table, f"{{{w}}}tr")
        nested_cell = ET.SubElement(nested_row, f"{{{w}}}tc")
        ET.SubElement(nested_cell, f"{{{w}}}tcPr")
        for text in ("NESTED_BEFORE", "NESTED_AFTER"):
            paragraph = ET.SubElement(nested_cell, f"{{{w}}}p")
            run = ET.SubElement(paragraph, f"{{{w}}}r")
            ET.SubElement(run, f"{{{w}}}t").text = text
            if text == "NESTED_BEFORE":
                nested_cell.append(altchunk)
        cell.append(nested_table)
    rels = ET.fromstring(members["word/_rels/document.xml.rels"])
    attributes = {
        "Id": "chunk", "Type": f"{r}/aFChunk",
        "Target": "https://example.invalid/no-fetch"
        if external else "chunks/content.html",
    }
    if external:
        attributes["TargetMode"] = "External"
    ET.SubElement(rels, f"{{{rels_ns}}}Relationship", attributes)
    if not external:
        members["word/chunks/content.html"] = (
            b"<html><body>ALTCHUNK_CELL_VISIBLE</body></html>")
        types = ET.fromstring(members["[Content_Types].xml"])
        ET.SubElement(types, f"{{{types_ns}}}Override", {
            "PartName": "/word/chunks/content.html", "ContentType": "text/html"})
        members["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8")
    members["word/document.xml"] = ET.tostring(document, encoding="utf-8")
    members["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("nested", [False, True])
@pytest.mark.parametrize("external", [False, True])
def test_docx_altchunk_in_any_table_cell_is_unprocessed(
        workflow, tmp_path, nested, external):
    path = tmp_path / f"altchunk-cell-{nested}-{external}.docx"
    _write_docx_altchunk_in_table(path, nested=nested, external=external)
    _, coverage = _project(workflow, [path])
    assert coverage["entries"][0]["status"] == "partial"
    assert any("altChunk" in item["error"]
               and "word/document.xml" in item["error"]
               for item in coverage["unprocessed"])


def test_readable_docx_headers_footers_and_tables_reach_discovery(tmp_path):
    from test_agent_host import session_at, submission

    path = tmp_path / "stories.docx"
    _write_docx_stories(path)
    request = submission(tmp_path)
    request["attachments"] = [str(path)]
    with session_at(tmp_path / "case") as session:
        run = session.prepare_submission(request)
        coverage = session.materials(run["run_id"])
        texts = [_text(session.workflow, item) for item in coverage["segments"]]
        assert texts == [
            "BODY_FIRST", "BODY_CELL", "BODY_LAST",
            "HEADER_VISIBLE", "HEADER_CELL", "FOOTER_VISIBLE", "FOOTER_CELL",
        ]
        header = next(item for item in coverage["segments"]
                      if _text(session.workflow, item) == "HEADER_VISIBLE")
        assert header["locator"] == {
            "member": "word/header1.xml", "story": "header", "paragraph": 1}
        assert coverage["entries"][0]["status"] == "projected"
        assert coverage["unprocessed"] == []
        assert [item["text"] for batch in session.discovery_batches(run["run_id"])
                for item in batch["items"]] == texts
    assert [item["text"] for item in intake.extract_docx_paragraphs(
        path.read_bytes()).locators] == ["BODY_FIRST", "BODY_LAST"]


@pytest.mark.parametrize("tag", ["sdt", "txbxContent", "ins", "del", "fldSimple",
                                  "drawing", "pict", "object"])
def test_docx_unsupported_visible_structures_are_not_reported_complete(workflow, tmp_path, tag):
    from docx.oxml import parse_xml

    document = docx.Document()
    document.add_paragraph("普通正文")
    document.element.body.insert(1, parse_xml(
        f'<w:{tag} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:t>未处理对象内的文字</w:t></w:r></w:p>'
        f'</w:{tag}>'))
    path = tmp_path / "content-control.docx"
    document.save(path)
    _, coverage = _project(workflow, [path])
    assert coverage["entries"][0]["status"] == "partial"
    assert any("word/document.xml" in item["error"] and tag in item["error"]
               for item in coverage["unprocessed"])


@pytest.mark.parametrize("story,member", [
    ("footnotes", "word/footnotes.xml"),
    ("footnotes", "word/notes/footnotes-custom.xml"),
    ("endnotes", "word/notes/endnotes-custom.xml"),
    ("comments", "word/notes/comments-custom.xml"),
])
def test_docx_additional_text_parts_are_explicitly_unprocessed(workflow, tmp_path, story, member):
    from xml.etree import ElementTree as ET

    path = tmp_path / "notes.docx"
    _write_docx_stories(path)
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    rels = ET.fromstring(members["word/_rels/document.xml.rels"])
    ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
        "Id": "rIdSyntheticFootnotes",
        "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{story}",
        "Target": member.removeprefix("word/"),
    })
    types = ET.fromstring(members["[Content_Types].xml"])
    ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
        "PartName": "/" + member,
        "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{story}+xml",
    })
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    kind = story[:-1]
    members[member] = (
        f'<w:{story} xmlns:w="{w}">'
        f'<w:{kind} w:id="1"><w:p><w:r><w:t>NOTE_VISIBLE</w:t></w:r></w:p>'
        f'</w:{kind}></w:{story}>').encode("utf-8")
    main = ET.fromstring(members["word/document.xml"])
    run = main.find(f".//{{{w}}}p/{{{w}}}r")
    ET.SubElement(run, f"{{{w}}}{kind}Reference", {f"{{{w}}}id": "1"})
    members["word/document.xml"] = ET.tostring(main, encoding="utf-8")
    members["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8")
    members["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    _, coverage = _project(workflow, [path])
    assert coverage["entries"][0]["status"] == "partial"
    assert any(member in item["error"] for item in coverage["unprocessed"])


@pytest.mark.parametrize("external", [False, True])
def test_docx_altchunk_is_located_without_fetching_external_content(workflow, tmp_path, external):
    from xml.etree import ElementTree as ET

    path = tmp_path / "altchunk.docx"
    _write_docx_stories(path)
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    document = ET.fromstring(members["word/document.xml"])
    body = document.find(f"{{{w}}}body")
    body.insert(len(body) - 1, ET.Element(f"{{{w}}}altChunk", {f"{{{r}}}id": "chunk"}))
    rels = ET.fromstring(members["word/_rels/document.xml.rels"])
    attributes = {"Id": "chunk", "Type": f"{r}/aFChunk",
                  "Target": "https://example.invalid/no-fetch" if external else "chunks/content.html"}
    if external:
        attributes["TargetMode"] = "External"
    ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                  attributes)
    members["word/document.xml"] = ET.tostring(document, encoding="utf-8")
    members["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8")
    if not external:
        members["word/chunks/content.html"] = b"<html><body>ALTCHUNK_VISIBLE</body></html>"
        types = ET.fromstring(members["[Content_Types].xml"])
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
            "PartName": "/word/chunks/content.html", "ContentType": "text/html"})
        members["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    _, coverage = _project(workflow, [path])
    assert coverage["entries"][0]["status"] == "partial"
    assert any("word/document.xml" in item["error"] and "altChunk" in item["error"]
               for item in coverage["unprocessed"])


@pytest.mark.parametrize("change", ["text", "omit", "locator", "source", "projection"])
def test_resealed_coverage_must_still_derive_from_frozen_originals(tmp_path, change):
    from test_agent_host import session_at, submission

    request = submission(tmp_path)
    path = Path(request["attachments"][0])
    path.write_text("A" * 32768 + "ORIGINAL_END", encoding="utf-8")
    other = tmp_path / "other.txt"
    other.write_text("OTHER_ORIGINAL", encoding="utf-8")
    request["attachments"].append(str(other))
    with session_at(tmp_path / "case") as session:
        run = session.prepare_submission(request)
        original = session.materials(run["run_id"])
        changed = copy.deepcopy(original)
        segment = changed["segments"][0]
        if change == "text":
            blob = session.blobs.put_bytes(("B" * segment["char_count"]).encode("utf-8"))
            segment.update(text_blob_sha256=blob.sha256, text_sha256=blob.sha256)
        elif change == "omit":
            changed["segments"].pop(0)
        elif change == "locator":
            segment["locator"]["start"] += 1
        elif change == "source":
            segment["source_id"] = changed["segments"][-1]["source_id"]
        else:
            segment["projection_id"] = changed["segments"][-1]["projection_id"]
        if change != "omit":
            segment["segment_id"] = "SEGMENT::" + digest({
                key: value for key, value in segment.items() if key != "segment_id"})
        changed["coverage_digest"] = intake.product_coverage_digest({
            key: value for key, value in changed.items() if key != "coverage_digest"})
        with session.store._conn:
            session.store._conn.execute(
                "DELETE FROM product_objects WHERE object_id=?", ("COVERAGE::" + run["run_id"],))
        intake._seal_product_coverage(session.workflow, changed)
        before = session.store._conn.total_changes
        for operation in (session.materials, session.discovery_batches):
            with pytest.raises((intake.IntakeRejected, ValueError), match="原件|投影|派生|覆盖|coverage"):
                operation(run["run_id"])
        assert session.store._conn.total_changes == before
        assert session.blobs.read_bytes(run["inputs"][0]["blob_sha256"]).startswith(b"A")


def test_text_line_offsets_preserve_utf8_newlines_and_empty_lines(workflow, tmp_path):
    path = tmp_path / "lines.txt"
    data = "甲\r\n\nbeta\r终".encode("utf-8")
    path.write_bytes(data)
    _, coverage = _project(workflow, [path])
    segments = coverage["segments"]
    assert [_text(workflow, s) for s in segments] == ["甲\r\n", "\n", "beta\r", "终"]
    assert [s["locator"] for s in segments] == [
        {"line": 1, "start": 0, "end": 5},
        {"line": 2, "start": 5, "end": 6},
        {"line": 3, "start": 6, "end": 11},
        {"line": 4, "start": 11, "end": 14},
    ]
    assert b"".join(data[s["locator"]["start"]:s["locator"]["end"]]
                    for s in segments) == data


@pytest.mark.parametrize("name,data,status", [
    ("empty.txt", b"", "empty"),
    ("invalid.txt", b"\xff\xfe", "failed"),
    ("broken.pdf", b"not pdf", "failed"),
    ("broken.docx", b"not docx", "failed"),
    ("unsupported.bin", b"some content", "unsupported"),
])
def test_unreadable_or_empty_inputs_remain_in_coverage(workflow, tmp_path, name, data, status):
    path = tmp_path / name
    path.write_bytes(data)
    _, coverage = _project(workflow, [path])
    assert coverage["received_count"] == len(coverage["entries"]) == 1
    assert coverage["entries"][0]["status"] == status
    assert coverage["entries"][0]["blob_sha256"] == sha256_hex(data)
    assert coverage["unprocessed"]
    assert coverage["segments"] == []


def test_zip_all_members_and_duplicate_bytes_remain_distinct(workflow, tmp_path):
    text = tmp_path / "outside.txt"
    text.write_bytes(b"same synthetic bytes\n")
    path = tmp_path / "materials.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("folder/", b"")
        archive.writestr("folder/one.txt", text.read_bytes())
        archive.writestr("folder/two.txt", text.read_bytes())
        archive.writestr("unknown.exe", b"binary")
        archive.writestr("nested.zip", b"PK not expanded")
    _, coverage = _project(workflow, [text, path])
    assert coverage["received_count"] == 2
    assert len(coverage["entries"]) == 7
    assert len(coverage["segments"]) == 3
    sources = [s["source_id"] for s in coverage["segments"]]
    assert len(set(sources)) == 3
    archive_entry = coverage["entries"][1]
    assert all(e["parent_id"] == archive_entry["material_id"]
               for e in coverage["entries"][2:])
    assert [e["status"] for e in coverage["entries"][2:]] == [
        "directory", "projected", "projected", "unsupported", "unsupported"]
    assert not (tmp_path / "folder").exists()
    assert any("嵌套" in u["error"] for u in coverage["unprocessed"])


@pytest.mark.parametrize("name", ["../escape.txt", "/abs.txt", "C:\\escape.txt",
                                  "\\\\server\\share.txt", "dir/../../bad.txt"])
def test_zip_unsafe_paths_keep_every_member_without_extracting(workflow, tmp_path, name):
    path = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("safe.txt", b"safe")
        archive.writestr(name, b"unsafe")
    _, coverage = _project(workflow, [path])
    assert len(coverage["entries"]) == 3
    assert all(e["status"] == "rejected" and e["error"] for e in coverage["entries"])
    assert coverage["segments"] == []
    assert len(coverage["unprocessed"]) == 3
    assert not (tmp_path / "safe.txt").exists()


def test_zip_duplicate_names_reject_both_occurrences(workflow, tmp_path):
    path = tmp_path / "duplicates.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("same.txt", "one")
        with pytest.warns(UserWarning, match="Duplicate name"):
            archive.writestr("same.txt", "two")
    _, coverage = _project(workflow, [path])
    assert len(coverage["entries"]) == 3
    assert len({e["material_id"] for e in coverage["entries"]}) == 3
    assert all(e["status"] == "rejected" for e in coverage["entries"])
    assert coverage["segments"] == []


def test_zip_crc_read_failure_keeps_member_and_usable_sibling(workflow, tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("bad.txt", b"corrupt-me")
        archive.writestr("good.txt", b"readable sibling")
    data = bytearray(buffer.getvalue())
    offset = data.index(b"corrupt-me")
    data[offset] ^= 1
    path = tmp_path / "crc.zip"
    path.write_bytes(data)
    _, coverage = _project(workflow, [path])
    assert len(coverage["entries"]) == 3
    assert coverage["entries"][1]["status"] == "failed"
    assert coverage["entries"][1]["blob_sha256"] is None
    assert coverage["entries"][1]["error"]
    assert [_text(workflow, s) for s in coverage["segments"]] == ["readable sibling"]


@pytest.mark.parametrize("kind", ["count", "total", "ratio"])
def test_zip_safety_limits_preserve_central_directory(workflow, tmp_path, kind):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        if kind == "count":
            for index in range(1001):
                archive.writestr(f"{index}.txt", b"x")
        else:
            archive.writestr("large.txt", b"\0" * (1024 * 1024) if kind == "ratio" else b"x")
    data = bytearray(buffer.getvalue())
    if kind == "total":
        central = data.index(b"PK\x01\x02")
        struct.pack_into("<I", data, central + 24, 256 * 1024 * 1024 + 1)
    path = tmp_path / "limit.zip"
    path.write_bytes(data)
    _, coverage = _project(workflow, [path])
    assert len(coverage["entries"]) == (1002 if kind == "count" else 2)
    assert all(e["status"] == "rejected" for e in coverage["entries"])
    assert coverage["segments"] == []


def test_file_byte_and_count_limits_keep_input_denominator(workflow, tmp_path):
    freeze, project = _api()
    path = tmp_path / "large.txt"
    with path.open("wb") as handle:
        handle.truncate(64 * 1024 * 1024 + 1)
    inputs = freeze([str(path)], workflow.blobs)
    assert inputs[0]["status"] == "rejected"
    assert inputs[0]["blob_sha256"] is None
    assert project(workflow, "RUN::large", inputs)["unprocessed"]
    many = freeze([str(path)] * 257, workflow.blobs)
    assert len(many) == 257
    assert all(item["status"] == "rejected" for item in many)
    assert project(workflow, "RUN::count", many)["received_count"] == 257


def test_total_input_limit_counts_each_occurrence(workflow, tmp_path):
    freeze, _ = _api()
    path = tmp_path / "large.bin"
    with path.open("wb") as handle:
        handle.truncate(64 * 1024 * 1024)
    inputs = freeze([str(path)] * 5, workflow.blobs)
    assert [i["status"] for i in inputs] == ["frozen"] * 4 + ["rejected"]
    assert inputs[-1]["blob_sha256"] is None
    assert inputs[-1]["error"]


def test_long_segment_is_losslessly_split_with_offsets(workflow, tmp_path):
    path = tmp_path / "long.txt"
    text = "长" * 32769 + "\r\n"
    path.write_bytes(text.encode("utf-8"))
    _, coverage = _project(workflow, [path])
    segments = coverage["segments"]
    assert len(segments) == 2
    assert max(s["char_count"] for s in segments) <= 32768
    assert "".join(_text(workflow, s) for s in segments) == text
    assert [(s["locator"]["char_start"], s["locator"]["char_end"])
            for s in segments] == [(0, 32768), (32768, 32771)]
    assert [(s["locator"]["start"], s["locator"]["end"])
            for s in segments] == [(0, 98304), (98304, 98309)]


def test_frozen_blob_corruption_stops_projection(workflow, tmp_path):
    freeze, project = _api()
    path = tmp_path / "original.txt"
    path.write_bytes(b"original")
    inputs = freeze([str(path)], workflow.blobs)
    blob = inputs[0]["blob_sha256"]
    (workflow.blobs.root / "objects" / blob[:2] / blob).write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="复核|身份|blob"):
        project(workflow, "RUN::tampered", inputs)
    assert workflow.store.list_product_objects("RUN::tampered", "coverage") == []


@pytest.mark.parametrize("action", ["change", "deny"])
def test_source_change_and_read_failure_keep_denominator(workflow, tmp_path, action):
    freeze, project = _api()
    path = tmp_path / "changing.txt"
    path.write_bytes(b"before")
    armed = [True]

    def at_read(event, args):
        if event == "open" and str(args[0]) == str(path) \
                and args[1] == "r" and armed[0]:
            armed[0] = False
            if action == "deny":
                raise PermissionError("合成读取失败")
            path.write_bytes(b"changed during read")

    sys.addaudithook(at_read)
    try:
        inputs = freeze([str(path)], workflow.blobs)
    finally:
        armed[0] = False
    assert inputs[0]["status"] == "failed"
    assert inputs[0]["blob_sha256"] is None
    assert inputs[0]["error"]
    coverage = project(workflow, "RUN::" + action, inputs)
    assert coverage["received_count"] == len(coverage["entries"]) == 1
    assert coverage["segments"] == []
    assert len(coverage["unprocessed"]) == 1


def test_corrupt_deflate_stream_retains_usable_sibling(workflow, tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("broken.txt", b"compressed synthetic data" * 20)
        archive.writestr("good.txt", b"usable")
    data = bytearray(buffer.getvalue())
    name_size, extra_size = struct.unpack_from("<HH", data, 26)
    data[30 + name_size + extra_size] = 0xff
    path = tmp_path / "broken-stream.zip"
    path.write_bytes(data)
    _, coverage = _project(workflow, [path])
    assert len(coverage["entries"]) == 3
    assert coverage["entries"][1]["status"] == "failed"
    assert coverage["entries"][1]["error"]
    assert [_text(workflow, s) for s in coverage["segments"]] == ["usable"]


@pytest.mark.parametrize("suffix,media", [
    (".txt", "text/plain"), (".md", "text/markdown"), (".zip", "application/zip"),
])
def test_product_source_media_matches_actual_parser(workflow, tmp_path, suffix, media):
    path = tmp_path / ("material" + suffix)
    if suffix == ".zip":
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("content.txt", "synthetic")
    else:
        path.write_bytes(b"synthetic")
    _, coverage = _project(workflow, [path])
    source = workflow.store.fetch_one(
        "sources", "source_id", coverage["entries"][0]["source_id"])
    assert source["media_type"] == media


def _canonical_json_bytes(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":"), allow_nan=False).encode("utf-8")


def test_product_coverage_digest_streams_large_local_identity_without_relaxing_host():
    coverage_digest = getattr(intake, "product_coverage_digest", None)
    assert callable(coverage_digest), "缺少本地完整覆盖流式摘要"
    small = {"segments": [{"text": "合成\n", "locator": {"line": 1}}],
             "unknown": None, "values": [True, 0, 1.5]}
    assert coverage_digest(small) == digest(small)
    large = {"unprocessed": [{"error": "测" * 1500000}]}
    assert coverage_digest(large) == sha256_hex(_canonical_json_bytes(large))
    with pytest.raises(ProductRejected, match="4MiB"):
        digest(large)
    with pytest.raises(ValueError):
        coverage_digest({"not_json": float("nan")})


def test_coverage_chunks_preserve_all_6000_lines_and_replay(workflow, tmp_path):
    path = tmp_path / "many-lines.txt"
    path.write_bytes(b"x\n" * 6000)
    run_id = "RUN::" + "a" * 64
    inputs, coverage = _project(workflow, [path], run_id)
    assert coverage["received_count"] == 1
    assert len(coverage["segments"]) == 6000
    assert coverage["unprocessed"] == []
    assert [s["locator"] for s in coverage["segments"]] == [
        {"line": line, "start": (line - 1) * 2, "end": line * 2}
        for line in range(1, 6001)]
    assert "".join(_text(workflow, s) for s in coverage["segments"]) == "x\n" * 6000
    body = {key: value for key, value in coverage.items() if key != "coverage_digest"}
    assert len(_canonical_json_bytes(body)) > 4 * 1024 * 1024
    assert coverage["coverage_digest"] == sha256_hex(_canonical_json_bytes(body))
    index = workflow.store.get_product_object(
        "COVERAGE::" + run_id, run_id=run_id, kind="coverage")
    assert index["schema_version"] == "ag1.coverage.index.v1"
    assert index["counts"] == {"entries": 1, "segments": 6000, "unprocessed": 0}
    assert not {"entries", "segments", "unprocessed"} & set(index)
    assert len(index["parts"]) >= 3
    assert index["index_digest"] == digest({
        key: value for key, value in index.items() if key != "index_digest"})
    assert len(_canonical_json_bytes(index)) < 4 * 1024 * 1024
    restored = {"entries": [], "segments": [], "unprocessed": []}
    for ordinal, ref in enumerate(index["parts"]):
        part = workflow.store.get_product_object(
            ref["part_id"], run_id=run_id, kind="coverage_part")
        assert part["schema_version"] == "ag1.coverage.part.v1"
        assert part["run_id"] == run_id
        assert part["inputs_digest"] == index["inputs_digest"]
        assert part["coverage_digest"] == index["coverage_digest"]
        assert part["part_index"] == ref["part_index"] == ordinal
        assert part["collection"] == ref["collection"]
        assert part["start"] == ref["start"] == len(restored[ref["collection"]])
        assert part["count"] == ref["count"] == len(part["items"])
        assert ref["byte_length"] == len(_canonical_json_bytes(part)) < 4 * 1024 * 1024
        assert ref["part_digest"] == part["part_digest"] == digest({
            key: value for key, value in part.items() if key != "part_digest"})
        restored[part["collection"]].extend(part["items"])
    assert restored == {key: coverage[key] for key in restored}
    path.unlink()
    assert intake.project_product_materials(workflow, run_id, inputs) == coverage
    assert len(workflow.store.list_product_objects(
        run_id, "coverage_part")) == len(index["parts"])


def _rewrite_synthetic_product_object(workflow, object_id, body, *, rehash=True):
    """合成Case的篡改注入；同步库行hash以单独验证分片引用闭包。"""
    if rehash:
        for field in ("index_digest", "part_digest"):
            if field in body:
                body[field] = digest({key: value for key, value in body.items()
                                      if key != field})
    saved = json.loads(workflow.store._conn.execute(
        "SELECT body_json FROM product_objects WHERE object_id=?", (object_id,)
    ).fetchone()[0])
    if saved.get("schema_version") == "ag1.product-record.v1":
        saved["body"] = body
    else:
        saved = body
    packed = json.dumps(saved, ensure_ascii=False, sort_keys=True, allow_nan=False)
    with workflow.store._conn:
        workflow.store._conn.execute(
            "UPDATE product_objects SET body_json=?,digest=? WHERE object_id=?",
            (packed, sha256_hex(packed.encode("utf-8")), object_id))


@pytest.mark.parametrize("change", [
    "deleted", "bytes", "ownership", "part_order", "part_count", "part_start",
    "part_collection", "reference_hash", "reference_order", "reference_duplicate",
    "collection_count", "received_count",
])
def test_coverage_chunks_reject_missing_or_changed_parts(workflow, tmp_path, change):
    path = tmp_path / "input.txt"
    path.write_bytes(b"one\ntwo\n")
    run_id = "RUN::chunks-integrity"
    inputs, _ = _project(workflow, [path, tmp_path / "missing.txt"], run_id)
    index_id = "COVERAGE::" + run_id
    index = workflow.store.get_product_object(index_id, run_id=run_id, kind="coverage")
    assert index["schema_version"] == "ag1.coverage.index.v1", "未启用分片封存"
    part_id = index["parts"][0]["part_id"]
    part = workflow.store.get_product_object(part_id, run_id=run_id, kind="coverage_part")
    if change == "deleted":
        with workflow.store._conn:
            workflow.store._conn.execute(
                "DELETE FROM product_objects WHERE object_id=?", (part_id,))
    elif change in {"reference_hash", "reference_order", "reference_duplicate",
                    "collection_count", "received_count"}:
        if change == "reference_hash":
            index["parts"][0]["part_digest"] = "0" * 64
        elif change == "reference_order":
            index["parts"][0], index["parts"][1] = index["parts"][1], index["parts"][0]
        elif change == "reference_duplicate":
            index["parts"].append(copy.deepcopy(index["parts"][0]))
        elif change == "collection_count":
            index["counts"]["segments"] += 1
        else:
            index["received_count"] += 1
        _rewrite_synthetic_product_object(workflow, index_id, index)
    else:
        if change == "bytes":
            part["items"][0]["locator"]["filename"] = "tampered.txt"
        elif change == "ownership":
            part["run_id"] = "RUN::other"
        elif change == "part_order":
            part["part_index"] += 1
        elif change == "part_count":
            part["count"] += 1
        elif change == "part_start":
            part["start"] += 1
        else:
            part["collection"] = "segments"
        _rewrite_synthetic_product_object(workflow, part_id, part)
        for field in ("part_index", "count", "start", "collection", "part_digest"):
            index["parts"][0][field] = part[field]
        index["parts"][0]["byte_length"] = len(_canonical_json_bytes(part))
        _rewrite_synthetic_product_object(workflow, index_id, index)
    with pytest.raises((ValueError, intake.IntakeRejected), match="分片|coverage|身份|归属"):
        intake.project_product_materials(workflow, run_id, inputs)
    if change == "deleted":
        assert workflow.store.get_product_object(part_id) is None


@pytest.mark.parametrize("corrupt_digest", [False, True])
def test_coverage_legacy_unchunked_objects_are_verified(workflow, tmp_path, corrupt_digest):
    path = tmp_path / "old.txt"
    path.write_bytes(b"old frozen text\n")
    run_id = "RUN::legacy"
    inputs, coverage = _project(workflow, [path], run_id)
    with workflow.store._conn:
        workflow.store._conn.execute(
            "DELETE FROM product_objects WHERE run_id=?", (run_id,))
    if corrupt_digest:
        coverage["coverage_digest"] = "0" * 64
    workflow.store.put_product_object(
        "COVERAGE::" + run_id, run_id=run_id, kind="coverage", body=coverage)
    path.unlink()
    if corrupt_digest:
        with pytest.raises(intake.IntakeRejected, match="coverage|摘要|身份"):
            intake.project_product_materials(workflow, run_id, inputs)
    else:
        assert intake.project_product_materials(workflow, run_id, inputs) == coverage
    assert workflow.store.list_product_objects(run_id, "coverage_part") == []


def test_coverage_chunks_publish_index_last_and_reuse_orphan_parts(workflow, tmp_path):
    path = tmp_path / "input.txt"
    path.write_bytes(b"synthetic")
    freeze, project = _api()
    inputs = freeze([str(path)], workflow.blobs)
    run_id = "RUN::interrupted"
    with workflow.store._conn:
        workflow.store._conn.execute("""
            CREATE TRIGGER fail_coverage_index BEFORE INSERT ON product_objects
            WHEN NEW.kind='coverage'
            BEGIN SELECT RAISE(ABORT, 'synthetic index interruption'); END
        """)
    with pytest.raises(sqlite3.IntegrityError, match="synthetic index interruption"):
        project(workflow, run_id, inputs)
    assert workflow.store.get_product_object("COVERAGE::" + run_id) is None
    parts = workflow.store.list_product_objects(run_id, "coverage_part")
    assert parts, "根索引必须晚于完整分片发布"
    with workflow.store._conn:
        workflow.store._conn.execute("DROP TRIGGER fail_coverage_index")
    coverage = project(workflow, run_id, inputs)
    assert workflow.store.list_product_objects(run_id, "coverage_part") == parts
    assert len(coverage["segments"]) == 1


@pytest.mark.parametrize("field", ["part_index", "start", "count"])
def test_coverage_chunks_reject_boolean_part_positions(workflow, tmp_path, field):
    path = tmp_path / "input.txt"
    path.write_bytes(b"synthetic")
    run_id = "RUN::boolean-position"
    inputs, _ = _project(workflow, [path], run_id)
    index_id = "COVERAGE::" + run_id
    index = workflow.store.get_product_object(index_id, run_id=run_id, kind="coverage")
    ref = index["parts"][0]
    part = workflow.store.get_product_object(
        ref["part_id"], run_id=run_id, kind="coverage_part")
    part[field] = bool(part[field])
    _rewrite_synthetic_product_object(workflow, ref["part_id"], part)
    ref["part_digest"] = part["part_digest"]
    ref["byte_length"] = len(_canonical_json_bytes(part))
    _rewrite_synthetic_product_object(workflow, index_id, index)
    with pytest.raises(intake.IntakeRejected, match="分片"):
        intake.project_product_materials(workflow, run_id, inputs)
